from io import BytesIO
import csv
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from app.database import get_db
from app.models import User, AssistantResult
from app.schemas import AssistantQueryRequest, AssistantQueryResponse
from app.security import get_current_user
from app.services.assistant_service import run_assistant_query
from app.services.audit_service import log_audit

router = APIRouter(prefix="/assistant", tags=["Natural Language Assistant"])

@router.post("/query", response_model=AssistantQueryResponse)
def query_assistant(
    req: AssistantQueryRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Feature 9: Natural-language assistant
    Plain-language questions (maturing FDRs, unmatched transactions, incomplete release requests, why a payment was flagged).
    Scoped strictly by user role and region. Refuses out-of-scope requests. Links to evidence.
    """
    res = run_assistant_query(db=db, user=current_user, query_text=req.query)

    # Record search query in audit log
    log_audit(
        db=db,
        actor_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action=f"Assistant Query: '{req.query}' (Intent: {res['intent_matched']}, Denied: {res['access_denied']})",
        entity_type="AssistantQuery",
        entity_id=current_user.id,
        after_state={"query": req.query, "intent": res["intent_matched"], "denied": res["access_denied"]}
    )
    db.commit()

    return res


@router.get("/results/{result_id}/export")
def export_assistant_result(
    result_id: str,
    format: str = "xlsx",
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    result = db.query(AssistantResult).filter(
        AssistantResult.id == result_id,
        AssistantResult.owner_id == current_user.id,
    ).first()
    if not result:
        raise HTTPException(status_code=404, detail="Assistant result not found.")
    if format not in {"xlsx", "csv"}:
        raise HTTPException(status_code=400, detail="Export format must be xlsx or csv.")

    payload = result.payload
    rows = payload.get("table", {}).get("rows", [])
    columns = payload.get("table", {}).get("columns", [])
    intent_code = payload.get("intent_code", "")
    filename_prefix = "AICTE_mismatch_cases" if intent_code == "MISMATCH_CASES" else "AICTE_assistant_result"
    filename = f"{filename_prefix}_{result.created_at.strftime('%Y%m%d_%H%M')}"
    log_audit(
        db=db,
        actor_id=current_user.id,
        actor_name=current_user.full_name,
        actor_role=current_user.role,
        action=f"Assistant export: {format}",
        entity_type="AssistantResult",
        entity_id=result.id,
        after_state={"format": format, "row_count": len(rows)},
    )
    db.commit()
    buffer = BytesIO()

    if format == "csv":
        text_buffer = BytesIO()
        content = [column.get("label", column.get("key", "")) for column in columns]
        lines = [content]
        lines.extend([[str(row.get(column.get("key", ""), "")) for column in columns] for row in rows])
        import io
        writer_buffer = io.StringIO()
        writer = csv.writer(writer_buffer)
        writer.writerows(lines)
        text_buffer.write(writer_buffer.getvalue().encode("utf-8-sig"))
        text_buffer.seek(0)
        return StreamingResponse(text_buffer, media_type="text/csv", headers={
            "Content-Disposition": f'attachment; filename="{filename}.csv"',
        })

    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Font, PatternFill

    def safe_cell_value(value):
        if not isinstance(value, str):
            return value
        if value[:1] in {"=", "+", "-", "@"}:
            return "'" + value
        return value

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["AICTE Assistant Result", ""])
    summary.append(["Question", result.question])
    summary.append(["Interpretation", payload.get("interpretation", {}).get("text", "")])
    summary.append(["Scope", result.scope_text])
    summary.append(["Generated at", result.created_at.isoformat()])
    summary.append(["Row count", len(rows)])
    summary.append(["Mode", result.mode])
    summary.append(["Disclaimer", "Software-generated from supplied records; not an official certification"])

    data = workbook.create_sheet("Data")
    keys = [column.get("key", "") for column in columns]
    data.append([column.get("label", key) for column, key in zip(columns, keys)])
    for cell in data[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163E63")
    for row in rows:
        data.append([safe_cell_value(row.get(key, "")) for key in keys])
    data.freeze_panes = "A2"
    data.auto_filter.ref = data.dimensions
    for index, column in enumerate(columns, start=1):
        if column.get("type") == "currency":
            for cell in data.iter_cols(min_col=index, max_col=index, min_row=2):
                for value_cell in cell:
                    value_cell.number_format = '₹#,##0.00'
    status_index = next((index for index, column in enumerate(columns, start=1) if column.get("key") == "status"), None)
    if status_index:
        status_letter = data.cell(row=1, column=status_index).column_letter
        data.conditional_formatting.add(
            f"{status_letter}2:{status_letter}{max(2, data.max_row)}",
            FormulaRule(formula=[f'ISNUMBER(SEARCH("mismatch",{status_letter}2))'], fill=PatternFill("solid", fgColor="FDE2E2")),
        )
    for column in data.columns:
        data.column_dimensions[column[0].column_letter].width = min(80, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
    details = workbook.create_sheet("Query details")
    details.append(["Intent", payload.get("intent_matched", "")])
    details.append(["Intent code", intent_code])
    details.append(["Scope", result.scope_text])
    details.append(["Result ID", result.id])
    charts = workbook.create_sheet("Charts")
    charts.append(["Chart", "Label", "Value"])
    for cell in charts[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="163E63")
    chart = payload.get("charts", [{}])[0] if payload.get("charts") else {}
    chart_rows = chart.get("series", [])
    for item in chart_rows:
        charts.append([chart.get("title", "Assistant chart"), item.get("label", ""), item.get("value", 0)])
    if chart_rows:
        native_chart = BarChart()
        native_chart.title = chart.get("title", "Assistant chart")
        native_chart.y_axis.title = "Value"
        native_chart.x_axis.title = "Label"
        native_chart.add_data(Reference(charts, min_col=3, min_row=1, max_row=charts.max_row), titles_from_data=True)
        native_chart.set_categories(Reference(charts, min_col=2, min_row=2, max_row=charts.max_row))
        charts.add_chart(native_chart, "E2")
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={
        "Content-Disposition": f'attachment; filename="{filename}.xlsx"',
    })
